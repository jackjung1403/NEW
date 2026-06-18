#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""FINAL MV pitch deck — embeds the 5 nano-banana keyframes.
Run once the image host is reachable; downloads images then builds the deck."""
import json, os, urllib.request
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BASE="https://d8j0ntlcm91z4.cloudfront.net/user_2wflbbXerazanvSmxvNfLdSxoAb/"
IMGS={
 "concept": BASE+"hf_20260618_065123_64747103-cef7-4981-8c67-1c1c5e2e2cfd.png",
 "verse":   BASE+"hf_20260618_065126_d7e0ea46-ed8e-4b6d-b401-2367f184808b.png",
 "chorus":  BASE+"hf_20260618_065127_ecd56ab8-f1cc-42c6-a4eb-09c6c1171032.png",
 "silence": BASE+"hf_20260618_065129_206eccd9-12ec-409f-8188-b89288a767ad.png",
 "finale":  BASE+"hf_20260618_065131_4e63fc69-c56d-49a2-a1b7-dcef0a487e03.png",
}
os.makedirs("deck_assets", exist_ok=True)
paths={}
for k,u in IMGS.items():
    p=f"deck_assets/{k}.png"
    if not os.path.exists(p) or os.path.getsize(p)<2000:
        urllib.request.urlretrieve(u,p)
    assert os.path.getsize(p)>2000, f"download failed for {k} ({os.path.getsize(p)}b) — host blocked?"
    paths[k]=p
print("images ready:", {k:os.path.getsize(v) for k,v in paths.items()})

BG=RGBColor(0x0D,0x0E,0x12); PANEL=RGBColor(0x17,0x1A,0x22); TX=RGBColor(0xEC,0xEE,0xF4)
MUT=RGBColor(0x9A,0xA0,0xB0); BLUE=RGBColor(0x5B,0x8C,0xFF); RED=RGBColor(0xFF,0x4D,0x5E)
WARM=RGBColor(0xFF,0xB2,0x4D); GREEN=RGBColor(0x3D,0xDC,0x84); LINE=RGBColor(0x2A,0x2E,0x3A)
FONT="Arial"; EMU_W,EMU_H=Inches(13.333),Inches(7.5)
data=json.load(open("pitch_data.json"))
prs=Presentation(); prs.slide_width=EMU_W; prs.slide_height=EMU_H
BLANK=prs.slide_layouts[6]

def slide(bg=BG):
    s=prs.slides.add_slide(BLANK)
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,EMU_W,EMU_H)
    r.fill.solid(); r.fill.fore_color.rgb=bg; r.line.fill.background(); r.shadow.inherit=False
    return s
def rect(s,x,y,w,h,fill=None,line=None,lw=1.0,rnd=False):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rnd else MSO_SHAPE.RECTANGLE,x,y,w,h)
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw)
    shp.shadow.inherit=False; return shp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sa=4,ls=1.0):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Emu(0)
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(sa); p.space_before=Pt(0); p.line_spacing=ls
        for (t,sz,col,b) in para:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.color.rgb=col; r.font.bold=b; r.font.name=FONT
    return tb
def pic(s,path,x,y,w,h,line=LINE):
    ph=s.shapes.add_picture(path,x,y,width=w,height=h)
    fr=rect(s,x,y,w,h,fill=None,line=line,lw=1.0); return ph
def tone_strip(s,y,h=Inches(0.12)):
    bw=EMU_W/3
    for i,c in enumerate([RGBColor(0x2F,0x4F,0x8F),RED,WARM]):
        rect(s,Emu(int(bw*i)),y,Emu(int(bw)+9525),h,fill=c)
def footer(s,n):
    txt(s,Inches(0.7),Inches(7.04),Inches(10),Inches(0.3),[[("一遍一遍 · ADD · MV PITCH",9,MUT,False)]])
    txt(s,Inches(12.2),Inches(7.04),Inches(0.9),Inches(0.3),[[(str(n),9,MUT,False)]],align=PP_ALIGN.RIGHT)
def header(s,kick,title,accent):
    rect(s,Inches(0.5),Inches(0.55),Inches(0.06),Inches(0.95),fill=accent)
    tone_strip(s,Inches(0.0))
    txt(s,Inches(0.7),Inches(0.55),Inches(11),Inches(0.4),[[(kick.upper(),12,accent,True)]])
    txt(s,Inches(0.7),Inches(0.92),Inches(12),Inches(0.8),[[(title,30,TX,True)]])

def content_slide(n,kick,title,accent,items,img=None):
    s=slide(); header(s,kick,title,accent)
    pw=Inches(6.7) if img else Inches(11.9)
    rect(s,Inches(0.7),Inches(2.0),pw,Inches(4.6),fill=PANEL,line=LINE,rnd=True)
    runs=[]
    for it in items:
        runs.append([("›  ",14,accent,True),(it,14,TX,False)]); runs.append([("",4,MUT,False)])
    txt(s,Inches(1.0),Inches(2.28),pw-Inches(0.6),Inches(4.1),runs,ls=1.08)
    if img:
        pic(s,paths[img],Inches(7.65),Inches(2.0),Inches(4.95),Inches(2.78))
        txt(s,Inches(7.65),Inches(4.95),Inches(4.95),Inches(1.4),
            [[("KEY VISUAL — nano-banana",11,MUT,True)]],ls=1.05)
    footer(s,n); return s
def twocol(s_n,kick,title,accent,comp):
    s=slide(); header(s,kick,title,accent)
    for head,its,col,xx in [("OUR ADVANTAGES",comp.get("our_advantages",[]),GREEN,0.7),
                            ("AVOIDING",comp.get("competitors",[]),MUT,6.75)]:
        rect(s,Inches(xx),Inches(2.0),Inches(5.85),Inches(4.6),fill=PANEL,line=col,rnd=True)
        txt(s,Inches(xx+0.3),Inches(2.25),Inches(5.3),Inches(0.4),[[(head,16,col,True)]])
        runs=[]
        for it in its:
            runs.append([("•  ",13,col,True),(it,13,TX,False)]); runs.append([("",4,MUT,False)])
        txt(s,Inches(xx+0.3),Inches(2.8),Inches(5.3),Inches(3.6),runs,ls=1.1)
    footer(s,s_n); return s

# 1 — TITLE (concept image full-bleed, dimmed)
s=slide()
s.shapes.add_picture(paths["concept"],0,0,width=EMU_W,height=EMU_H)
ov=rect(s,0,0,EMU_W,EMU_H,fill=RGBColor(0x06,0x07,0x0A));
ov.fill.fore_color.rgb=RGBColor(0x06,0x07,0x0A)
# transparency
from pptx.oxml.ns import qn
sp=ov.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
a=sp.makeelement(qn('a:alpha'),{'val':'42000'}); sp.append(a)
tone_strip(s,Inches(0.0)); tone_strip(s,Inches(7.32))
txt(s,Inches(0.7),Inches(2.5),Inches(12),Inches(0.5),[[("MUSIC VIDEO · PITCH",14,RGBColor(0xD8,0xDB,0xE3),True)]])
txt(s,Inches(0.7),Inches(2.95),Inches(12),Inches(1.6),[[(data["company_name"],46,RGBColor(0xFF,0xFF,0xFF),True)]])
txt(s,Inches(0.72),Inches(4.2),Inches(12),Inches(0.6),[[(data.get("tagline",""),22,WARM,True)]])
txt(s,Inches(0.7),Inches(6.45),Inches(12),Inches(0.5),
    [[("Beijing · Ultra-low-budget · One-location one-take · Sunset",13,RGBColor(0xD0,0xD3,0xDD),False)]])

# sections with images mapped
secmap=[("problem","The Brief","과제 — 무엇을 풀어야 하나",RED,None),
        ("solution","The Concept","콘셉트 — 반복되는 도시",WARM,"concept"),
        ("market","Artist & Market","ADD & 시장 맥락",BLUE,None),
        ("product","Production","연출 — 단일 로케 원테이크",BLUE,"verse"),
        ("traction","Key Shots & Hooks","핵심 샷 & 훅",RED,"chorus"),
        ("business_model","Budget Strategy","예산 전략",GREEN,None),
        ("team","Team","ADD 6인 + 크루",BLUE,"finale"),
        ("financials","Next Steps","다음 스텝",WARM,None)]
n=2
for key,kick,title,accent,img in secmap:
    if key in data: content_slide(n,kick,title,accent,data[key],img=img); n+=1

# KEYFRAMES grid slide (2x2): verse/chorus/silence(무음)/finale
s=slide(); header(s,"Keyframes","톤 키프레임 — nano-banana",WARM)
grid=[("VERSE · 블루 · 작가의 밤","verse"),("CHORUS · 레드 · 6인 군무","chorus"),
      ("무음 / Leavin' · 실루엣","silence"),("FINALE · 웜 · 선셋","finale")]
gw,gh=Inches(5.85),Inches(2.05)
for i,(cap,key) in enumerate(grid):
    xx=Inches(0.7)+(i%2)*Inches(6.05); yy=Inches(2.0)+(i//2)*Inches(2.35)
    pic(s,paths[key],xx,yy,gw,gh)
    txt(s,xx+Inches(0.1),yy+gh+Inches(0.02),gw,Inches(0.3),[[(cap,12,MUT,True)]])
footer(s,n); n+=1

# EDGE two-col
if "competition" in data and isinstance(data["competition"],dict):
    twocol(n,"The Edge","경쟁 우위",WARM,data["competition"]); n+=1

# CLOSING (finale image dimmed)
s=slide()
s.shapes.add_picture(paths["finale"],0,0,width=EMU_W,height=EMU_H)
ov=rect(s,0,0,EMU_W,EMU_H); ov.fill.solid(); ov.fill.fore_color.rgb=RGBColor(0x06,0x07,0x0A)
sp=ov.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
sp.append(sp.makeelement(qn('a:alpha'),{'val':'45000'}))
tone_strip(s,Inches(7.32))
txt(s,Inches(0.7),Inches(2.7),Inches(12),Inches(1.4),
    [[(data["company_name"].split("—")[0].strip(),52,RGBColor(0xFF,0xFF,0xFF),True)]])
txt(s,Inches(0.72),Inches(3.95),Inches(12),Inches(0.6),[[(data.get("tagline",""),20,WARM,True)]])
txt(s,Inches(0.7),Inches(5.0),Inches(12),Inches(0.5),
    [[("ADD · Beijing · One-location performance film",14,RGBColor(0xD0,0xD3,0xDD),False)]])

out="yibian_ADD_pitch_deck_FINAL.pptx"
prs.save(out); print("saved",out,len(prs.slides._sldIdLst),"slides")
